"""Evidence-first resume review: never invents experience."""
import json,re
from collections import Counter
from pathlib import Path
STOP=set("and the to for with from that this are you your our their have has will can into using use these those a an of on in at by as or is be we it job work team skills experience required preferred ability strong knowledge".split())
def read_resume(path):
    p=Path(path).expanduser()
    if not p.is_file():raise FileNotFoundError(str(p))
    ext=p.suffix.lower()
    if ext in (".txt",".md"):return p.read_text(encoding="utf8")
    if ext==".pdf":
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(p)).pages)
    if ext==".docx":
        from docx import Document
        doc=Document(str(p))
        parts=[x.text for x in doc.paragraphs]
        parts.extend(" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows)
        return "\n".join(parts)
    raise ValueError("Resume must be .pdf, .docx, .txt, or .md")
def tokens(s):
    return [w for w in re.findall(r"[a-z][a-z0-9+#.\-]*",s.lower()) if len(w)>2 and w not in STOP]
def generate_tailored_markdown(resume,job_description,role="Target position"):
    if not resume.strip() or not job_description.strip():raise ValueError("Resume and job description must contain text")
    demands=Counter(tokens(job_description)); evidence=set(tokens(resume))
    ranked=[w for w,_ in demands.most_common(35)]
    found=[w for w in ranked if w in evidence][:20]
    missing=[w for w in ranked if w not in evidence][:20]
    bullets=[line.strip() for line in resume.splitlines() if line.strip() and (line.lstrip().startswith(("-", "•", "*")) or len(line.strip())>55)]
    relevant=sorted(bullets,key=lambda line:sum(demands.get(w,0) for w in set(tokens(line))),reverse=True)[:6]
    report={"role":role,"verified_terms":found,"not_verified_in_resume":missing,"relevant_existing_bullets":relevant,"notice":"Only existing resume content is evidenced. Verify and approve all edits."}
    body=["# Resume tailoring review: "+role,"","## Verified skills and terminology",", ".join(found) if found else "No obvious matches found.","","## Requirements not verified in the resume","Review, do not add without supporting experience: "+(", ".join(missing) if missing else "None identified"),"","## Existing experience to consider highlighting"]
    body.extend("- "+line for line in relevant)
    body.extend(["","## Original resume (unchanged)","",resume])
    return "\n".join(body)+"\n",report
def export_tailoring(markdown,path):
    p=Path(path).expanduser();p.parent.mkdir(parents=True,exist_ok=True)
    if p.suffix.lower() in (".md",".txt"):
        p.write_text(markdown,encoding="utf8")
    elif p.suffix.lower()==".docx":
        from docx import Document
        d=Document()
        for line in markdown.splitlines():d.add_paragraph(line)
        d.save(p)
    elif p.suffix.lower()==".pdf":
        from reportlab.pdfgen import canvas
        from reportlab.lib.pagesizes import letter
        c=canvas.Canvas(str(p),pagesize=letter);y=740
        for line in markdown.splitlines():
            for i in range(0,max(1,len(line)),95):
                if y<55:c.showPage();y=740
                c.drawString(36,y,line[i:i+95].encode("latin1",errors="replace").decode("latin1"));y-=14
        c.save()
    else:raise ValueError("Output must be .md, .txt, .pdf or .docx")
    return p
